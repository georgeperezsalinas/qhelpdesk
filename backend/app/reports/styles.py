"""
styles.py – Estilos ERP compartidos para Excel y PDF
Paleta institucional: azul oscuro #1B3A6B, acento #2F7FD1, gris #F5F6FA
"""
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ── COLORES ───────────────────────────────────────────────────────────────────
COLOR_PRIMARY   = "1B3A6B"   # azul oscuro encabezados
COLOR_ACCENT    = "2F7FD1"   # azul medio subencabezados
COLOR_LIGHT_BG  = "EBF3FB"   # fondo filas alternas
COLOR_GRAY      = "F5F6FA"   # fondo sección
COLOR_WHITE     = "FFFFFF"
COLOR_SUCCESS   = "1E7E34"
COLOR_WARNING   = "856404"
COLOR_DANGER    = "721C24"
COLOR_GOLD      = "856404"

# ── FILLS ─────────────────────────────────────────────────────────────────────
FILL_PRIMARY  = PatternFill("solid", fgColor=COLOR_PRIMARY)
FILL_ACCENT   = PatternFill("solid", fgColor=COLOR_ACCENT)
FILL_LIGHT    = PatternFill("solid", fgColor=COLOR_LIGHT_BG)
FILL_GRAY     = PatternFill("solid", fgColor=COLOR_GRAY)
FILL_WHITE    = PatternFill("solid", fgColor=COLOR_WHITE)
FILL_SUCCESS  = PatternFill("solid", fgColor="D4EDDA")
FILL_WARNING  = PatternFill("solid", fgColor="FFF3CD")
FILL_DANGER   = PatternFill("solid", fgColor="F8D7DA")

# ── FONTS ─────────────────────────────────────────────────────────────────────
FONT_TITLE    = Font(name="Calibri", bold=True, size=16, color=COLOR_WHITE)
FONT_SUBTITLE = Font(name="Calibri", bold=True, size=12, color=COLOR_WHITE)
FONT_HEADER   = Font(name="Calibri", bold=True, size=10, color=COLOR_WHITE)
FONT_SUBHEAD  = Font(name="Calibri", bold=True, size=10, color=COLOR_PRIMARY)
FONT_BODY     = Font(name="Calibri", size=10)
FONT_SMALL    = Font(name="Calibri", size=9, color="595959")
FONT_BOLD     = Font(name="Calibri", bold=True, size=10)
FONT_SUCCESS  = Font(name="Calibri", bold=True, size=10, color=COLOR_SUCCESS)
FONT_WARNING  = Font(name="Calibri", bold=True, size=10, color=COLOR_WARNING)
FONT_DANGER   = Font(name="Calibri", bold=True, size=10, color=COLOR_DANGER)

# ── ALIGNMENT ─────────────────────────────────────────────────────────────────
ALIGN_CENTER  = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT    = Alignment(horizontal="left",   vertical="center", wrap_text=True)
ALIGN_RIGHT   = Alignment(horizontal="right",  vertical="center")

# ── BORDERS ───────────────────────────────────────────────────────────────────
THIN  = Side(style="thin",   color="D9D9D9")
MED   = Side(style="medium", color="BFBFBF")
BORDER_THIN   = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=Side(style="medium", color=COLOR_ACCENT))

def auto_width(ws, min_w=8, max_w=50):
    """Ajusta el ancho de todas las columnas automáticamente."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value else ""
                max_len = max(max_len, len(val))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = max(min_w, min(max_w, max_len + 2))

def freeze_header(ws, row=2):
    ws.freeze_panes = ws.cell(row=row, column=1)

def add_filters(ws, row=1):
    ws.auto_filter.ref = ws.dimensions
