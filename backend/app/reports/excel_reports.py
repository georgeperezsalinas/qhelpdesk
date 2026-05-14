"""
excel_reports.py – Generadores de reportes Excel ERP nivel empresarial
"""
import io
from datetime import datetime, date, timedelta
from typing import Optional
from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from sqlalchemy.orm import Session
from sqlalchemy import func

from app.reports.styles import *
from app.models.ticket import Ticket, EstadoTicket, PrioridadTicket
from app.models.inventario import Equipo, Licencia, EstadoEquipo
from app.models.mantenimiento import OrdenMantenimiento, TipoMantenimiento, EstadoOrden
from app.models.compra import SolicitudCompra, EstadoSolicitud
from app.models.usuario import Usuario, RolUsuario
from app.models.backup import EjecucionBackup, EstadoBackup


# ─────────────────────────────────────────────────────────────────────────────
def _header_sheet(ws, titulo: str, subtitulo: str, entidad: str = "Entidad Pública"):
    """Genera la cabecera estándar ERP en la hoja."""
    ws.row_dimensions[1].height = 45
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 18

    # Fila 1 – Título principal
    ws.merge_cells("A1:N1")
    c = ws["A1"]
    c.value       = f"  QHELP DESK ERP  ·  {titulo.upper()}"
    c.font        = FONT_TITLE
    c.fill        = FILL_PRIMARY
    c.alignment   = ALIGN_LEFT

    # Fila 2 – Subtítulo
    ws.merge_cells("A2:N2")
    c = ws["A2"]
    c.value       = f"  {subtitulo}"
    c.font        = FONT_SUBTITLE
    c.fill        = FILL_ACCENT
    c.alignment   = ALIGN_LEFT

    # Fila 3 – Metadatos
    ws.merge_cells("A3:G3")
    ws["A3"].value     = f"  {entidad}"
    ws["A3"].font      = FONT_SMALL
    ws["A3"].alignment = ALIGN_LEFT

    ws.merge_cells("H3:N3")
    ws["H3"].value     = f"  Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  QHELP DESK ERP v1.0"
    ws["H3"].font      = FONT_SMALL
    ws["H3"].alignment = Alignment(horizontal="right", vertical="center")

    ws.row_dimensions[5].height = 8   # espacio


def _header_row(ws, row: int, cols: list):
    """Genera fila de encabezado de tabla."""
    ws.row_dimensions[row].height = 22
    for i, col_name in enumerate(cols, 1):
        c = ws.cell(row=row, column=i, value=col_name)
        c.font      = FONT_HEADER
        c.fill      = FILL_PRIMARY
        c.alignment = ALIGN_CENTER
        c.border    = BORDER_THIN


def _data_row(ws, row: int, values: list, alternate: bool = False):
    fill = FILL_LIGHT if alternate else FILL_WHITE
    ws.row_dimensions[row].height = 18
    for i, val in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=val)
        c.font      = FONT_BODY
        c.fill      = fill
        c.alignment = ALIGN_LEFT
        c.border    = BORDER_THIN


def _kpi_box(ws, row: int, col: int, label: str, value, color_fill=None, color_font=None):
    """Mini card de KPI en una celda combinada."""
    end_col = col + 2
    ws.merge_cells(
        start_row=row, start_column=col,
        end_row=row+1, end_column=end_col
    )
    c = ws.cell(row=row, column=col)
    c.value = f"{label}\n{value}"
    c.font  = Font(name="Calibri", bold=True, size=11,
                   color=color_font or COLOR_PRIMARY)
    c.fill  = PatternFill("solid", fgColor=color_fill or "EBF3FB") if color_fill else FILL_LIGHT
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(
        left=Side(style="medium", color=COLOR_ACCENT),
        right=Side(style="medium", color=COLOR_ACCENT),
        top=Side(style="medium", color=COLOR_ACCENT),
        bottom=Side(style="medium", color=COLOR_ACCENT),
    )


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE 1: TICKETS — Gestión y SLA
# ══════════════════════════════════════════════════════════════════════════════
def reporte_tickets(
    db: Session,
    fecha_desde: Optional[date] = None,
    fecha_hasta: Optional[date] = None,
) -> bytes:
    if not fecha_desde: fecha_desde = date.today().replace(day=1)
    if not fecha_hasta: fecha_hasta = date.today()

    wb = Workbook()

    # ── Hoja 1: Resumen ejecutivo ─────────────────────────────────────────────
    ws = wb.active
    ws.title = "Resumen Ejecutivo"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3

    _header_sheet(ws, "Gestión de Tickets",
                  f"Período: {fecha_desde.strftime('%d/%m/%Y')} – {fecha_hasta.strftime('%d/%m/%Y')}")

    # KPIs
    q = db.query(Ticket).filter(
        Ticket.creado_en >= datetime.combine(fecha_desde, datetime.min.time()),
        Ticket.creado_en <= datetime.combine(fecha_hasta, datetime.max.time()),
    )
    total        = q.count()
    resueltos    = q.filter(Ticket.estado.in_([EstadoTicket.resuelto, EstadoTicket.cerrado])).count()
    vencidos     = q.filter(
        Ticket.sla_limite < datetime.utcnow(),
        Ticket.estado.notin_([EstadoTicket.resuelto, EstadoTicket.cerrado, EstadoTicket.cancelado])
    ).count()
    sla_ok       = q.filter(Ticket.sla_cumplido == True).count()
    tasa_sla     = f"{round(sla_ok/max(resueltos,1)*100,1)}%"
    t_prom_res   = db.query(func.avg(Ticket.tiempo_resolucion_h)).filter(
        Ticket.tiempo_resolucion_h != None,
        Ticket.creado_en >= datetime.combine(fecha_desde, datetime.min.time()),
    ).scalar()
    nps_prom     = db.query(func.avg(Ticket.nps_puntuacion)).filter(
        Ticket.nps_puntuacion != None
    ).scalar()

    row = 6
    kpis = [
        ("Total tickets",    total,                         "EBF3FB", COLOR_PRIMARY),
        ("Resueltos",        resueltos,                     "D4EDDA", COLOR_SUCCESS),
        ("SLA cumplido",     tasa_sla,                      "D4EDDA", COLOR_SUCCESS),
        ("SLA vencidos",     vencidos,                      "F8D7DA" if vencidos > 0 else "D4EDDA",
                                                             COLOR_DANGER if vencidos > 0 else COLOR_SUCCESS),
        ("Tiempo prom (h)",  f"{round(t_prom_res or 0, 1)}h", "EBF3FB", COLOR_PRIMARY),
        ("NPS promedio",     f"{round(nps_prom or 0, 1)}/10",  "FFF3CD", COLOR_GOLD),
    ]
    col = 2
    for label, val, cf, ff in kpis:
        _kpi_box(ws, row, col, label, val, cf, ff)
        col += 4
    row += 3

    # Por prioridad
    ws.cell(row=row, column=2, value="DISTRIBUCIÓN POR PRIORIDAD").font = FONT_SUBHEAD
    row += 1
    _header_row(ws, row, ["Prioridad", "Total", "Resueltos", "Pendientes", "SLA Cumplido", "% SLA"])
    row += 1
    for prioridad in [PrioridadTicket.critica, PrioridadTicket.alta,
                      PrioridadTicket.media, PrioridadTicket.baja]:
        t  = q.filter(Ticket.prioridad == prioridad).count()
        r  = q.filter(Ticket.prioridad == prioridad,
                      Ticket.estado.in_([EstadoTicket.resuelto, EstadoTicket.cerrado])).count()
        p  = t - r
        sk = q.filter(Ticket.prioridad == prioridad, Ticket.sla_cumplido == True).count()
        pct= f"{round(sk/max(r,1)*100,1)}%"
        _data_row(ws, row, [prioridad.value.capitalize(), t, r, p, sk, pct], row % 2 == 0)
        # colorear prioridad
        c = ws.cell(row=row, column=1)
        if prioridad == PrioridadTicket.critica: c.fill = FILL_DANGER;  c.font = FONT_DANGER
        elif prioridad == PrioridadTicket.alta:  c.fill = FILL_WARNING; c.font = FONT_WARNING
        row += 1
    row += 2

    # Por categoría
    ws.cell(row=row, column=2, value="DISTRIBUCIÓN POR CATEGORÍA").font = FONT_SUBHEAD
    row += 1
    _header_row(ws, row, ["Categoría", "Total", "% del total"])
    row += 1
    cats = db.query(Ticket.categoria, func.count(Ticket.id)).group_by(Ticket.categoria)\
              .order_by(func.count(Ticket.id).desc()).all()
    for cat, cnt in cats:
        pct = f"{round(cnt/max(total,1)*100,1)}%"
        _data_row(ws, row, [str(cat.value).capitalize() if cat else "Sin categoría", cnt, pct], row%2==0)
        row += 1
    row += 2

    # Por técnico
    ws.cell(row=row, column=2, value="DESEMPEÑO POR TÉCNICO").font = FONT_SUBHEAD
    row += 1
    _header_row(ws, row, ["Técnico", "Asignados", "Resueltos", "Pendientes", "SLA Cumplido", "T. Prom. (h)"])
    row += 1
    tecnicos = db.query(Usuario).filter(
        Usuario.rol.in_([RolUsuario.especialista, RolUsuario.mesa_ayuda])
    ).all()
    for tec in tecnicos:
        asig = q.filter(Ticket.tecnico_id == tec.id).count()
        if asig == 0: continue
        res  = q.filter(Ticket.tecnico_id == tec.id,
                        Ticket.estado.in_([EstadoTicket.resuelto, EstadoTicket.cerrado])).count()
        pend = asig - res
        slak = q.filter(Ticket.tecnico_id == tec.id, Ticket.sla_cumplido == True).count()
        tprom= db.query(func.avg(Ticket.tiempo_resolucion_h)).filter(
            Ticket.tecnico_id == tec.id, Ticket.tiempo_resolucion_h != None
        ).scalar()
        _data_row(ws, row, [
            f"{tec.nombre} {tec.apellido}", asig, res, pend, slak,
            f"{round(tprom or 0, 1)}h"
        ], row%2==0)
        row += 1

    auto_width(ws)
    freeze_header(ws, 7)

    # ── Hoja 2: Detalle completo ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Detalle de Tickets")
    ws2.sheet_view.showGridLines = False
    _header_sheet(ws2, "Detalle de Tickets",
                  f"Período: {fecha_desde.strftime('%d/%m/%Y')} – {fecha_hasta.strftime('%d/%m/%Y')}")

    cols = ["N° Ticket", "Título", "Prioridad", "Estado", "Categoría",
            "Canal", "Solicitante", "Técnico", "Sede",
            "Creado", "Resuelto", "SLA Límite", "SLA Cumplido",
            "T. Resolución (h)", "NPS"]
    _header_row(ws2, 6, cols)

    tickets = q.order_by(Ticket.creado_en.desc()).limit(5000).all()
    for i, t in enumerate(tickets, 7):
        sol = f"{t.solicitante.nombre} {t.solicitante.apellido}" if t.solicitante else "—"
        tec = f"{t.tecnico.nombre} {t.tecnico.apellido}" if t.tecnico else "Sin asignar"
        row_data = [
            t.numero, t.titulo,
            t.prioridad.value.capitalize(), t.estado.value.replace("_"," ").capitalize(),
            t.categoria.value.capitalize() if t.categoria else "—",
            t.canal_entrada.value if t.canal_entrada else "—",
            sol, tec,
            t.sede.nombre if t.sede else "—",
            t.creado_en.strftime("%d/%m/%Y %H:%M") if t.creado_en else "—",
            t.resuelto_en.strftime("%d/%m/%Y %H:%M") if t.resuelto_en else "—",
            t.sla_limite.strftime("%d/%m/%Y %H:%M") if t.sla_limite else "—",
            "✓ Sí" if t.sla_cumplido else ("✗ No" if t.sla_cumplido is False else "—"),
            t.tiempo_resolucion_h or "—",
            t.nps_puntuacion or "—",
        ]
        _data_row(ws2, i, row_data, i%2==0)
        # Colorear prioridad
        c = ws2.cell(row=i, column=3)
        if t.prioridad == PrioridadTicket.critica: c.fill=FILL_DANGER; c.font=FONT_DANGER
        elif t.prioridad == PrioridadTicket.alta:  c.fill=FILL_WARNING; c.font=FONT_WARNING
        # Colorear SLA
        c2 = ws2.cell(row=i, column=13)
        if t.sla_cumplido is True:  c2.fill=FILL_SUCCESS; c2.font=FONT_SUCCESS
        elif t.sla_cumplido is False: c2.fill=FILL_DANGER; c2.font=FONT_DANGER

    auto_width(ws2, max_w=40)
    freeze_header(ws2, 7)
    add_filters(ws2)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE 2: INVENTARIO TI — Activos y Depreciación
# ══════════════════════════════════════════════════════════════════════════════
def reporte_inventario(db: Session) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen Inventario"
    ws.sheet_view.showGridLines = False

    _header_sheet(ws, "Inventario TI", "Activos tecnológicos – Estado actual y depreciación")

    # KPIs
    total_eq  = db.query(func.count(Equipo.id)).scalar()
    activos   = db.query(func.count(Equipo.id)).filter(Equipo.estado == EstadoEquipo.activo).scalar()
    val_total = db.query(func.sum(Equipo.valor_compra)).filter(Equipo.estado != EstadoEquipo.de_baja).scalar() or 0
    gar_venc  = db.query(func.count(Equipo.id)).filter(
        Equipo.garantia_hasta < date.today(), Equipo.estado == EstadoEquipo.activo
    ).scalar()

    row = 6
    kpis = [
        ("Total equipos",    total_eq,                        "EBF3FB", COLOR_PRIMARY),
        ("Activos",          activos,                         "D4EDDA", COLOR_SUCCESS),
        ("Valor inventario", f"S/ {val_total:,.2f}",          "EBF3FB", COLOR_PRIMARY),
        ("Garantía vencida", gar_venc,
         "F8D7DA" if gar_venc > 0 else "D4EDDA",
         COLOR_DANGER if gar_venc > 0 else COLOR_SUCCESS),
    ]
    col = 2
    for label, val, cf, ff in kpis:
        _kpi_box(ws, row, col, label, val, cf, ff)
        col += 4
    row += 3

    # Por tipo
    ws.cell(row=row, column=2, value="DISTRIBUCIÓN POR TIPO DE EQUIPO").font = FONT_SUBHEAD
    row += 1
    _header_row(ws, row, ["Tipo", "Total", "Activos", "Mantenimiento", "Baja", "Valor Total (S/)"])
    row += 1
    tipos = db.query(Equipo.tipo, func.count(Equipo.id)).group_by(Equipo.tipo)\
               .order_by(func.count(Equipo.id).desc()).all()
    for tipo, cnt in tipos:
        act  = db.query(func.count(Equipo.id)).filter(Equipo.tipo==tipo, Equipo.estado==EstadoEquipo.activo).scalar()
        mant = db.query(func.count(Equipo.id)).filter(Equipo.tipo==tipo, Equipo.estado==EstadoEquipo.en_mantenimiento).scalar()
        baja = db.query(func.count(Equipo.id)).filter(Equipo.tipo==tipo, Equipo.estado==EstadoEquipo.de_baja).scalar()
        val  = db.query(func.sum(Equipo.valor_compra)).filter(Equipo.tipo==tipo).scalar() or 0
        _data_row(ws, row, [tipo.value.replace("_"," ").title(), cnt, act, mant, baja, f"S/ {val:,.2f}"], row%2==0)
        row += 1
    row += 2

    # Licencias por vencer
    ws.cell(row=row, column=2, value="LICENCIAS POR VENCER (próximos 60 días)").font = FONT_SUBHEAD
    row += 1
    _header_row(ws, row, ["Software", "Tipo", "Cantidad", "Vencimiento", "Días restantes", "Valor (S/)"])
    row += 1
    limite = date.today() + timedelta(days=60)
    lics = db.query(Licencia).filter(
        Licencia.fecha_vencimiento != None,
        Licencia.fecha_vencimiento <= limite,
        Licencia.activa == True,
    ).order_by(Licencia.fecha_vencimiento).all()
    for lic in lics:
        dias = (lic.fecha_vencimiento - date.today()).days
        _data_row(ws, row, [
            lic.software, lic.tipo_licencia, lic.cantidad_total,
            lic.fecha_vencimiento.strftime("%d/%m/%Y"),
            f"{dias} días" if dias >= 0 else "VENCIDA",
            f"S/ {lic.valor:,.2f}" if lic.valor else "—",
        ], row%2==0)
        c = ws.cell(row=row, column=5)
        if dias < 0:     c.fill=FILL_DANGER;  c.font=FONT_DANGER
        elif dias <= 15: c.fill=FILL_WARNING; c.font=FONT_WARNING
        row += 1

    auto_width(ws)
    freeze_header(ws, 7)

    # ── Hoja 2: Detalle equipos con depreciación ──────────────────────────────
    ws2 = wb.create_sheet("Detalle con Depreciación")
    ws2.sheet_view.showGridLines = False
    _header_sheet(ws2, "Inventario TI", "Detalle de equipos con cálculo de depreciación")

    cols = ["Código", "Código Patrimonial", "Tipo", "Marca", "Modelo", "Serie",
            "Estado", "Usuario Asignado", "Sede", "Ubicación",
            "F. Compra", "Valor Compra (S/)", "Vida Útil",
            "Valor Actual (S/)", "Dep. Anual (S/)", "% Depreciado",
            "Garantía hasta", "Garantía Vigente", "SO", "IP"]
    _header_row(ws2, 6, cols)

    equipos = db.query(Equipo).order_by(Equipo.codigo_inventario).all()
    for i, eq in enumerate(equipos, 7):
        # Calcular depreciación
        if eq.valor_compra and eq.fecha_compra:
            anios = (date.today() - eq.fecha_compra).days / 365.25
            dep_anual = (eq.valor_compra - (eq.valor_residual or 0)) / (eq.vida_util_anios or 4)
            val_act   = max(eq.valor_residual or 0, eq.valor_compra - dep_anual * anios)
            pct_dep   = min(100, round((1 - val_act/eq.valor_compra)*100, 1))
        else:
            dep_anual = val_act = pct_dep = "—"

        usr = eq.usuario_asignado
        _data_row(ws2, i, [
            eq.codigo_inventario, eq.codigo_patrimonial or "—",
            eq.tipo.value.replace("_"," ").title(), eq.marca, eq.modelo, eq.serie or "—",
            eq.estado.value.replace("_"," ").title(),
            f"{usr.nombre} {usr.apellido}" if usr else "Sin asignar",
            eq.sede.nombre if eq.sede else "—",
            eq.ubicacion_fisica or "—",
            eq.fecha_compra.strftime("%d/%m/%Y") if eq.fecha_compra else "—",
            f"S/ {eq.valor_compra:,.2f}" if eq.valor_compra else "—",
            f"{eq.vida_util_anios} años",
            f"S/ {val_act:,.2f}" if isinstance(val_act, float) else "—",
            f"S/ {dep_anual:,.2f}" if isinstance(dep_anual, float) else "—",
            f"{pct_dep}%" if isinstance(pct_dep, float) else "—",
            eq.garantia_hasta.strftime("%d/%m/%Y") if eq.garantia_hasta else "—",
            "Vigente" if eq.garantia_hasta and eq.garantia_hasta >= date.today() else ("Vencida" if eq.garantia_hasta else "—"),
            eq.sistema_operativo or "—",
            eq.ip_asignada or "—",
        ], i%2==0)

        # Colorear % depreciado
        if isinstance(pct_dep, float):
            c = ws2.cell(row=i, column=16)
            if pct_dep >= 80:   c.fill=FILL_DANGER;  c.font=FONT_DANGER
            elif pct_dep >= 50: c.fill=FILL_WARNING; c.font=FONT_WARNING

    auto_width(ws2, max_w=35)
    freeze_header(ws2, 7)
    add_filters(ws2)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE 3: MANTENIMIENTO
# ══════════════════════════════════════════════════════════════════════════════
def reporte_mantenimiento(db: Session,
                          fecha_desde: Optional[date] = None,
                          fecha_hasta: Optional[date] = None) -> bytes:
    if not fecha_desde: fecha_desde = date.today().replace(day=1)
    if not fecha_hasta: fecha_hasta = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen Mantenimiento"
    ws.sheet_view.showGridLines = False

    _header_sheet(ws, "Mantenimiento TI",
                  f"Período: {fecha_desde.strftime('%d/%m/%Y')} – {fecha_hasta.strftime('%d/%m/%Y')}")

    q = db.query(OrdenMantenimiento).filter(
        OrdenMantenimiento.fecha_programada >= fecha_desde,
        OrdenMantenimiento.fecha_programada <= fecha_hasta,
    )
    total    = q.count()
    prev     = q.filter(OrdenMantenimiento.tipo == TipoMantenimiento.preventivo).count()
    corr     = q.filter(OrdenMantenimiento.tipo == TipoMantenimiento.correctivo).count()
    comp     = q.filter(OrdenMantenimiento.estado == EstadoOrden.completado).count()
    costo_t  = db.query(func.sum(OrdenMantenimiento.costo)).filter(
        OrdenMantenimiento.fecha_programada >= fecha_desde,
        OrdenMantenimiento.estado == EstadoOrden.completado,
    ).scalar() or 0

    row = 6
    kpis = [
        ("Total órdenes", total,                   "EBF3FB", COLOR_PRIMARY),
        ("Preventivos",   prev,                    "D4EDDA", COLOR_SUCCESS),
        ("Correctivos",   corr,                    "FFF3CD", COLOR_GOLD),
        ("Completados",   comp,                    "D4EDDA", COLOR_SUCCESS),
        ("Costo total",   f"S/ {costo_t:,.2f}",   "EBF3FB", COLOR_PRIMARY),
        ("% Preventivo",  f"{round(prev/max(total,1)*100,1)}%", "D4EDDA", COLOR_SUCCESS),
    ]
    col = 2
    for label, val, cf, ff in kpis:
        _kpi_box(ws, row, col, label, val, cf, ff)
        col += 4
    row += 3

    # Detalle
    ws.cell(row=row, column=2, value="DETALLE DE ÓRDENES").font = FONT_SUBHEAD
    row += 1
    _header_row(ws, row, ["N° Orden", "Tipo", "Estado", "Equipo", "Técnico",
                            "F. Programada", "F. Completado", "Duración (min)",
                            "Costo (S/)", "Origen"])
    row += 1
    ordenes = q.order_by(OrdenMantenimiento.fecha_programada).all()
    for i, o in enumerate(ordenes):
        eq  = o.equipo
        tec = o.tecnico
        _data_row(ws, row+i, [
            o.numero,
            o.tipo.value.capitalize(),
            o.estado.value.replace("_"," ").capitalize(),
            f"{eq.marca} {eq.modelo} ({eq.codigo_inventario})" if eq else "—",
            f"{tec.nombre} {tec.apellido}" if tec else "—",
            o.fecha_programada.strftime("%d/%m/%Y"),
            o.fecha_fin.strftime("%d/%m/%Y %H:%M") if o.fecha_fin else "—",
            o.duracion_minutos or "—",
            f"S/ {o.costo:,.2f}" if o.costo else "—",
            o.origen.value if o.origen else "—",
        ], (row+i)%2==0)

    auto_width(ws)
    freeze_header(ws, 7)
    add_filters(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE 4: COMPRAS TI
# ══════════════════════════════════════════════════════════════════════════════
def reporte_compras(db: Session,
                    fecha_desde: Optional[date] = None,
                    fecha_hasta: Optional[date] = None) -> bytes:
    if not fecha_desde: fecha_desde = date.today().replace(day=1)
    if not fecha_hasta: fecha_hasta = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = "Compras TI"
    ws.sheet_view.showGridLines = False
    _header_sheet(ws, "Compras TI",
                  f"Período: {fecha_desde.strftime('%d/%m/%Y')} – {fecha_hasta.strftime('%d/%m/%Y')}")

    q = db.query(SolicitudCompra).filter(
        SolicitudCompra.creado_en >= datetime.combine(fecha_desde, datetime.min.time()),
        SolicitudCompra.creado_en <= datetime.combine(fecha_hasta, datetime.max.time()),
    )
    total     = q.count()
    aprobadas = q.filter(SolicitudCompra.estado.in_([
        EstadoSolicitud.aprobada, EstadoSolicitud.en_proceso, EstadoSolicitud.recibida
    ])).count()
    val_apro  = db.query(func.sum(SolicitudCompra.valor_aprobado)).filter(
        SolicitudCompra.estado.in_([EstadoSolicitud.aprobada, EstadoSolicitud.en_proceso, EstadoSolicitud.recibida])
    ).scalar() or 0
    val_recib = db.query(func.sum(SolicitudCompra.valor_final)).filter(
        SolicitudCompra.estado == EstadoSolicitud.recibida
    ).scalar() or 0

    row = 6
    kpis = [
        ("Total solicitudes", total,                   "EBF3FB", COLOR_PRIMARY),
        ("Aprobadas",         aprobadas,               "D4EDDA", COLOR_SUCCESS),
        ("Val. aprobado",     f"S/ {val_apro:,.2f}",  "EBF3FB", COLOR_PRIMARY),
        ("Val. ejecutado",    f"S/ {val_recib:,.2f}", "D4EDDA", COLOR_SUCCESS),
    ]
    col = 2
    for label, val, cf, ff in kpis:
        _kpi_box(ws, row, col, label, val, cf, ff)
        col += 4
    row += 3

    _header_row(ws, row, ["N° Solicitud", "Tipo", "Descripción", "Estado",
                            "Solicitante", "Aprobador", "Proveedor",
                            "Val. Estimado (S/)", "Val. Aprobado (S/)",
                            "F. Creación", "F. Aprobación", "F. Recepción",
                            "Presupuesto"])
    row += 1
    for i, s in enumerate(q.order_by(SolicitudCompra.creado_en.desc()).all()):
        sol = s.solicitante
        apr = s.aprobador
        prov= s.proveedor
        _data_row(ws, row+i, [
            s.numero, s.tipo.value.replace("_"," ").title() if s.tipo else "—",
            s.descripcion[:80],
            s.estado.value.replace("_"," ").title(),
            f"{sol.nombre} {sol.apellido}" if sol else "—",
            f"{apr.nombre} {apr.apellido}" if apr else "—",
            prov.razon_social if prov else "—",
            f"S/ {s.valor_estimado:,.2f}" if s.valor_estimado else "—",
            f"S/ {s.valor_aprobado:,.2f}" if s.valor_aprobado else "—",
            s.creado_en.strftime("%d/%m/%Y") if s.creado_en else "—",
            s.fecha_aprobacion.strftime("%d/%m/%Y") if s.fecha_aprobacion else "—",
            s.fecha_recepcion.strftime("%d/%m/%Y") if s.fecha_recepcion else "—",
            s.presupuesto_codigo or "—",
        ], (row+i)%2==0)
        c = ws.cell(row=row+i, column=4)
        if s.estado == EstadoSolicitud.rechazada: c.fill=FILL_DANGER; c.font=FONT_DANGER
        elif s.estado == EstadoSolicitud.recibida: c.fill=FILL_SUCCESS; c.font=FONT_SUCCESS

    auto_width(ws, max_w=40)
    freeze_header(ws, 7)
    add_filters(ws)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# REPORTE 5: GERENCIAL CONSOLIDADO
# ══════════════════════════════════════════════════════════════════════════════
def reporte_gerencial(db: Session,
                      fecha_desde: Optional[date] = None,
                      fecha_hasta: Optional[date] = None) -> bytes:
    if not fecha_desde: fecha_desde = date.today().replace(day=1)
    if not fecha_hasta: fecha_hasta = date.today()

    wb = Workbook()
    ws = wb.active
    ws.title = "Informe Gerencial"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3

    _header_sheet(ws, "Informe Gerencial – Oficina de Sistemas",
                  f"Período: {fecha_desde.strftime('%d/%m/%Y')} – {fecha_hasta.strftime('%d/%m/%Y')}")

    row = 6
    # ── Sección 1: Tickets ───────────────────────────────────────────────────
    ws.merge_cells(f"B{row}:N{row}")
    c = ws.cell(row=row, column=2, value="1.  GESTIÓN DE TICKETS DE SOPORTE")
    c.font = Font(name="Calibri", bold=True, size=12, color=COLOR_PRIMARY)
    c.fill = FILL_GRAY
    c.alignment = ALIGN_LEFT
    row += 1

    q = db.query(Ticket).filter(
        Ticket.creado_en >= datetime.combine(fecha_desde, datetime.min.time()),
    )
    total_t  = q.count()
    resuel_t = q.filter(Ticket.estado.in_([EstadoTicket.resuelto, EstadoTicket.cerrado])).count()
    sla_ok_t = q.filter(Ticket.sla_cumplido == True).count()
    nps_t    = db.query(func.avg(Ticket.nps_puntuacion)).filter(Ticket.nps_puntuacion!=None).scalar()

    for label, val in [
        ("Total tickets recibidos:", total_t),
        ("Tickets resueltos:", resuel_t),
        ("SLA cumplido:", f"{round(sla_ok_t/max(resuel_t,1)*100,1)}%"),
        ("NPS promedio:", f"{round(nps_t or 0,1)}/10"),
    ]:
        ws.cell(row=row, column=3, value=label).font = FONT_BODY
        c = ws.cell(row=row, column=5, value=val)
        c.font = FONT_BOLD
        c.alignment = ALIGN_RIGHT
        row += 1
    row += 1

    # ── Sección 2: Inventario ────────────────────────────────────────────────
    ws.merge_cells(f"B{row}:N{row}")
    c = ws.cell(row=row, column=2, value="2.  INVENTARIO Y ACTIVOS TI")
    c.font = Font(name="Calibri", bold=True, size=12, color=COLOR_PRIMARY)
    c.fill = FILL_GRAY
    c.alignment = ALIGN_LEFT
    row += 1

    total_eq  = db.query(func.count(Equipo.id)).scalar()
    val_inv   = db.query(func.sum(Equipo.valor_compra)).filter(Equipo.estado!=EstadoEquipo.de_baja).scalar() or 0
    lics_venc = db.query(func.count(Licencia.id)).filter(
        Licencia.fecha_vencimiento < date.today(), Licencia.activa==True).scalar()

    for label, val in [
        ("Total equipos activos:", db.query(func.count(Equipo.id)).filter(Equipo.estado==EstadoEquipo.activo).scalar()),
        ("Valor del inventario:", f"S/ {val_inv:,.2f}"),
        ("Licencias vencidas:", lics_venc),
    ]:
        ws.cell(row=row, column=3, value=label).font = FONT_BODY
        c = ws.cell(row=row, column=5, value=val)
        c.font = FONT_DANGER if (label=="Licencias vencidas:" and lics_venc>0) else FONT_BOLD
        c.alignment = ALIGN_RIGHT
        row += 1
    row += 1

    # ── Sección 3: Mantenimiento ─────────────────────────────────────────────
    ws.merge_cells(f"B{row}:N{row}")
    c = ws.cell(row=row, column=2, value="3.  MANTENIMIENTO")
    c.font = Font(name="Calibri", bold=True, size=12, color=COLOR_PRIMARY)
    c.fill = FILL_GRAY
    c.alignment = ALIGN_LEFT
    row += 1

    qm = db.query(OrdenMantenimiento).filter(
        OrdenMantenimiento.fecha_programada >= fecha_desde,
        OrdenMantenimiento.fecha_programada <= fecha_hasta,
    )
    prev_m = qm.filter(OrdenMantenimiento.tipo==TipoMantenimiento.preventivo).count()
    corr_m = qm.filter(OrdenMantenimiento.tipo==TipoMantenimiento.correctivo).count()
    pct_prev = round(prev_m/max(prev_m+corr_m,1)*100,1)

    for label, val in [
        ("Mantenimientos preventivos:", prev_m),
        ("Mantenimientos correctivos:", corr_m),
        ("% Preventivo (meta ≥60%):", f"{pct_prev}%"),
    ]:
        ws.cell(row=row, column=3, value=label).font = FONT_BODY
        c = ws.cell(row=row, column=5, value=val)
        c.font = FONT_SUCCESS if (label.startswith("%") and pct_prev>=60) else \
                 FONT_DANGER  if (label.startswith("%") and pct_prev<60) else FONT_BOLD
        c.alignment = ALIGN_RIGHT
        row += 1
    row += 1

    # ── Sección 4: Backup ────────────────────────────────────────────────────
    ws.merge_cells(f"B{row}:N{row}")
    c = ws.cell(row=row, column=2, value="4.  BACKUP Y CONTINUIDAD")
    c.font = Font(name="Calibri", bold=True, size=12, color=COLOR_PRIMARY)
    c.fill = FILL_GRAY
    c.alignment = ALIGN_LEFT
    row += 1

    hace7 = datetime.utcnow() - timedelta(days=7)
    total_bk = db.query(func.count(EjecucionBackup.id)).filter(EjecucionBackup.inicio>=hace7).scalar()
    exit_bk  = db.query(func.count(EjecucionBackup.id)).filter(
        EjecucionBackup.inicio>=hace7, EjecucionBackup.estado==EstadoBackup.exitoso).scalar()
    tasa_bk  = round(exit_bk/max(total_bk,1)*100,1)

    for label, val in [
        ("Backups ejecutados (7d):", total_bk),
        ("Backups exitosos (7d):", exit_bk),
        ("Tasa de éxito (meta ≥95%):", f"{tasa_bk}%"),
    ]:
        ws.cell(row=row, column=3, value=label).font = FONT_BODY
        c = ws.cell(row=row, column=5, value=val)
        c.font = FONT_SUCCESS if (label.startswith("Tasa") and tasa_bk>=95) else \
                 FONT_DANGER  if (label.startswith("Tasa") and tasa_bk<95) else FONT_BOLD
        c.alignment = ALIGN_RIGHT
        row += 1
    row += 2

    # ── Firma ────────────────────────────────────────────────────────────────
    ws.merge_cells(f"B{row}:N{row}")
    ws.cell(row=row, column=2,
            value=f"Informe generado automáticamente por QHELP DESK ERP · {datetime.now().strftime('%d/%m/%Y %H:%M')}").font = FONT_SMALL

    auto_width(ws)
    ws.column_dimensions["A"].width = 3

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
